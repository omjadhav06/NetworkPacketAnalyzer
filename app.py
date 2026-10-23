import csv
import io
import json
import html as html_lib
import subprocess
import threading
import string
import time
from collections import Counter, deque
from pathlib import Path

from flask import Flask, jsonify, render_template, Response, request
from scapy.all import AsyncSniffer, conf, get_if_list, rdpcap, IP, IPv6, TCP, UDP, ICMP, DNS

app = Flask(__name__)
BASE_DIR = Path(__file__).resolve().parent
CAPTURE_FILE = BASE_DIR / "captures" / "dns.cap"
MAX_LIVE_PACKETS = 1000
WIFI_INTERFACE = r"\Device\NPF_{BE2FE6F1-B2A2-4735-915E-F28695D21825}"

capture_lock = threading.RLock()
live_packets = deque(maxlen=MAX_LIVE_PACKETS)
live_packet_objects = {}  # packet number -> Scapy packet, retained for detail inspection
live_protocol_counts = Counter()
live_total_bytes = 0
live_total_seen = 0
live_sniffer = None
live_interface = None
live_error = None


def get_windows_adapters():
    """Return Windows adapter metadata indexed by GUID, if PowerShell is available."""
    if __import__("platform").system() != "Windows":
        return {}
    command = (
        "Get-NetAdapter -IncludeHidden | "
        "Select-Object Name,InterfaceDescription,Status,MacAddress,ifIndex,InterfaceGuid "
        "| ConvertTo-Json -Compress"
    )
    try:
        result = subprocess.run(
            ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", command],
            capture_output=True, text=True, timeout=5, check=True,
        )
        data = json.loads(result.stdout or "[]")
        if isinstance(data, dict):
            data = [data]
        adapters = {}
        for adapter in data:
            guid = str(adapter.get("InterfaceGuid") or "").strip("{} ").lower()
            if guid:
                adapters[guid] = adapter
        return adapters
    except (OSError, subprocess.SubprocessError, json.JSONDecodeError):
        app.logger.warning("Could not retrieve Windows adapter metadata", exc_info=True)
        return {}


def get_network_interfaces():
    """List Scapy interfaces with Windows names/status where available."""
    windows_adapters = get_windows_adapters()
    interfaces = []
    for iface in conf.ifaces.values():
        name = getattr(iface, "name", "") or ""
        description = getattr(iface, "description", "") or name
        if not name:
            continue
        adapter = None
        if name == r"\Device\NPF_Loopback":
            description = "Loopback (this computer only)"
            status = "Available"
            windows_name = "Loopback"
        else:
            guid = ""
            if "NPF_{" in name.upper():
                try:
                    guid = name.split("NPF_{", 1)[1].split("}", 1)[0].lower()
                except IndexError:
                    guid = ""
            adapter = windows_adapters.get(guid)
            if adapter:
                description = adapter.get("InterfaceDescription") or description
                status = str(adapter.get("Status") or "Unknown")
                windows_name = str(adapter.get("Name") or description)
            else:
                status = "Unknown"
                windows_name = description
        interfaces.append({
            "name": name,
            "description": description,
            "windows_name": windows_name,
            "status": status,
            "mac_address": (adapter or {}).get("MacAddress") or "",
            "if_index": (adapter or {}).get("ifIndex"),
            "recommended": name == WIFI_INTERFACE,
        })

    if not interfaces:
        interfaces = [{
            "name": name, "description": name, "windows_name": name,
            "status": "Unknown", "mac_address": "", "if_index": None,
            "recommended": name == WIFI_INTERFACE,
        } for name in get_if_list()]

    # Wi-Fi first, then connected/available adapters, then remaining adapters.
    def sort_key(item):
        status = item["status"].lower()
        if item["recommended"]:
            priority = 0
        elif item["name"] == r"\Device\NPF_Loopback":
            priority = 2
        elif status == "up":
            priority = 1
        elif status == "unknown":
            priority = 3
        else:
            priority = 4
        return priority, item["windows_name"].lower(), item["name"].lower()

    interfaces.sort(key=sort_key)
    return interfaces


def packet_details(packet, number):
    source_ip = destination_ip = "N/A"
    source_port = destination_port = "N/A"
    if IP in packet:
        source_ip, destination_ip = packet[IP].src, packet[IP].dst
    elif IPv6 in packet:
        source_ip, destination_ip = packet[IPv6].src, packet[IPv6].dst

    if TCP in packet:
        protocol = "TCP"
        source_port, destination_port = packet[TCP].sport, packet[TCP].dport
    elif UDP in packet:
        protocol = "DNS" if DNS in packet else "UDP"
        source_port, destination_port = packet[UDP].sport, packet[UDP].dport
    elif ICMP in packet:
        protocol = "ICMP"
    elif DNS in packet:
        protocol = "DNS"
    elif IPv6 in packet:
        protocol = "IPv6"
    elif IP in packet:
        protocol = "IP"
    else:
        protocol = packet.lastlayer().name

    tcp_flags = str(packet[TCP].flags) if TCP in packet else ""
    tcp_syn_only = bool(TCP in packet and (int(packet[TCP].flags) & 0x02) and not (int(packet[TCP].flags) & 0x10))
    return {
        "number": number,
        "timestamp": float(packet.time),
        "source_ip": source_ip,
        "destination_ip": destination_ip,
        "protocol": protocol,
        "source_port": source_port,
        "destination_port": destination_port,
        "tcp_flags": tcp_flags,
        "tcp_syn": tcp_syn_only,
        "size_bytes": len(packet),
    }



def _json_safe(value):
    """Convert Scapy field values into JSON-safe, readable values."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, bytes):
        printable = "".join(chr(byte) if chr(byte) in string.printable and byte not in (10, 13, 9) else "." for byte in value)
        return {"type": "bytes", "length": len(value), "hex": value.hex(" "), "ascii": printable}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    # Scapy enums, flags, addresses and other field wrappers are best displayed as text.
    try:
        return str(value)
    except Exception:
        return repr(value)


def _hex_dump(raw_bytes, width=16):
    """Create a classic offset / hexadecimal / ASCII packet dump."""
    lines = []
    for offset in range(0, len(raw_bytes), width):
        chunk = raw_bytes[offset:offset + width]
        hex_part = " ".join(f"{byte:02x}" for byte in chunk)
        ascii_part = "".join(chr(byte) if 32 <= byte <= 126 else "." for byte in chunk)
        lines.append({"offset": f"{offset:04x}", "hex": hex_part, "ascii": ascii_part})
    return lines


def inspect_packet(packet, number, source):
    """Return all decoded Scapy layers/fields plus a raw hexadecimal dump."""
    raw_bytes = bytes(packet)
    layers = []
    for layer_class in packet.layers():
        try:
            layer = packet.getlayer(layer_class)
            fields = {}
            for field_name, value in getattr(layer, "fields", {}).items():
                try:
                    fields[str(field_name)] = _json_safe(value)
                except Exception as error:
                    fields[str(field_name)] = f"<could not decode: {error}>"
            try:
                layer_bytes = bytes(layer)
            except Exception:
                layer_bytes = b""
            layers.append({
                "name": getattr(layer_class, "__name__", str(layer_class)),
                "summary": str(layer.summary()) if layer is not None else "",
                "fields": fields,
                "length_bytes": len(layer_bytes),
                "raw_hex": layer_bytes.hex(" "),
            })
        except Exception as error:
            layers.append({"name": str(layer_class), "summary": "Layer decode error", "fields": {"error": str(error)}, "length_bytes": 0, "raw_hex": ""})

    details = packet_details(packet, number)
    return {
        "source": source,
        "number": number,
        "timestamp": details["timestamp"],
        "timestamp_local": __import__("datetime").datetime.fromtimestamp(details["timestamp"]).astimezone().isoformat(timespec="milliseconds"),
        "summary": str(packet.summary()),
        "packet_length": len(raw_bytes),
        "metadata": details,
        "layers": layers,
        "raw_hex": raw_bytes.hex(" "),
        "hex_dump": _hex_dump(raw_bytes),
    }


def analyze_offline_pcap():
    if not CAPTURE_FILE.exists():
        raise FileNotFoundError(f"Capture file not found: {CAPTURE_FILE}")
    packets = rdpcap(str(CAPTURE_FILE))
    results = []
    protocol_counts = Counter()
    total_bytes = 0
    for number, packet in enumerate(packets, start=1):
        details = packet_details(packet, number)
        results.append(details)
        protocol_counts[details["protocol"]] += 1
        total_bytes += details["size_bytes"]
    summary = {
        "total_packets": len(results),
        "total_bytes": total_bytes,
        "protocol_count": len(protocol_counts),
        "average_packet_size": round(total_bytes / len(results), 2) if results else 0,
        "largest_packet": max((p["size_bytes"] for p in results), default=0),
    }
    return results, protocol_counts, summary


def live_packet_callback(packet):
    global live_total_bytes, live_total_seen
    try:
        with capture_lock:
            number = live_total_seen + 1
            details = packet_details(packet, number)
            live_total_seen += 1
            live_total_bytes += details["size_bytes"]
            live_protocol_counts[details["protocol"]] += 1
            live_packets.append(details)
            live_packet_objects[number] = packet
            # Keep detailed packet objects bounded just like the visible packet history.
            while len(live_packet_objects) > MAX_LIVE_PACKETS:
                del live_packet_objects[min(live_packet_objects)]
    except Exception as error:
        app.logger.exception("Could not process captured packet: %s", error)


def build_live_analytics(packet_history, window_seconds=30, max_items=8):
    """Build traffic-rate and top-talkers statistics from retained live packets."""
    now = time.time()
    current_second = int(now)
    bucket_seconds = list(range(current_second - window_seconds + 1, current_second + 1))
    packet_buckets = Counter()
    byte_buckets = Counter()
    source_ips = Counter()
    destination_ips = Counter()
    ports = Counter()

    for packet in packet_history:
        timestamp = float(packet.get("timestamp", now))
        second = int(timestamp)
        if current_second - window_seconds < second <= current_second:
            packet_buckets[second] += 1
            byte_buckets[second] += int(packet.get("size_bytes", 0))

        source_ip = packet.get("source_ip")
        destination_ip = packet.get("destination_ip")
        if source_ip and source_ip != "N/A":
            source_ips[source_ip] += 1
        if destination_ip and destination_ip != "N/A":
            destination_ips[destination_ip] += 1

        protocol = str(packet.get("protocol", "Other"))
        for direction in ("source_port", "destination_port"):
            port = packet.get(direction)
            if port not in (None, "N/A", ""):
                try:
                    ports[f"{port}/{protocol}"] += 1
                except TypeError:
                    pass

    return {
        "window_seconds": window_seconds,
        "labels": [time.strftime("%H:%M:%S", time.localtime(second)) for second in bucket_seconds],
        "packets_per_second": [packet_buckets[second] for second in bucket_seconds],
        "bytes_per_second": [byte_buckets[second] for second in bucket_seconds],
        "top_source_ips": [{"label": ip, "count": count} for ip, count in source_ips.most_common(max_items)],
        "top_destination_ips": [{"label": ip, "count": count} for ip, count in destination_ips.most_common(max_items)],
        "top_ports": [{"label": label, "count": count} for label, count in ports.most_common(max_items)],
        "history_packets": len(packet_history),
    }


def analyze_security_alerts(packet_history, window_seconds=30):
    """Generate heuristic alerts from recent live traffic; these are not proof of an attack."""
    now = time.time()
    recent = [p for p in packet_history if now - float(p.get("timestamp", now)) <= window_seconds]
    alerts = []

    # Traffic-rate spike: deliberately conservative default threshold, easy to tune.
    per_second = Counter(int(float(p.get("timestamp", now))) for p in recent)
    peak_rate = max(per_second.values(), default=0)
    if peak_rate >= 100:
        alerts.append({
            "id": "traffic-spike", "category": "Traffic spike", "severity": "warning", "title": "High traffic rate",
            "source_ip": "", "destination_ip": "", "destination_port": "", "observed_count": peak_rate,
            "threshold": 100, "window_seconds": 1,
            "message": f"Observed {peak_rate} packets in one second. This may be a burst of normal traffic or unusual activity.",
            "timestamp": now,
        })

    # TCP SYN scan heuristic: one source attempting many distinct destination ports.
    syn_by_source = {}
    syn_by_target = Counter()
    for packet in recent:
        if not packet.get("tcp_syn") or packet.get("source_ip") in (None, "N/A"):
            continue
        source = packet["source_ip"]
        destination = packet.get("destination_ip", "N/A")
        try:
            destination_port = int(packet.get("destination_port"))
        except (TypeError, ValueError):
            continue
        syn_by_source.setdefault(source, set()).add((destination, destination_port))
        syn_by_target[(source, destination, destination_port)] += 1

    for source, targets in syn_by_source.items():
        distinct_ports = {port for _destination, port in targets}
        if len(distinct_ports) >= 12:
            alerts.append({
                "id": f"port-scan-{source}", "category": "Possible port scan", "severity": "warning", "title": "Possible port scan pattern",
                "source_ip": source, "destination_ip": ", ".join(sorted({destination for destination, _port in targets})),
                "destination_port": ", ".join(str(port) for port in sorted(distinct_ports)),
                "observed_count": len(distinct_ports), "threshold": 12, "window_seconds": window_seconds,
                "message": f"{source} sent TCP SYN packets to {len(distinct_ports)} distinct destination ports in the last {window_seconds} seconds.",
                "timestamp": now,
            })

    # Repeated SYN attempts to one service may indicate retries or connection probing.
    watched_ports = {21: "FTP", 22: "SSH", 23: "Telnet", 445: "SMB", 1433: "MS SQL", 3306: "MySQL", 3389: "RDP", 4444: "common test/backdoor port", 5900: "VNC", 6379: "Redis", 27017: "MongoDB"}
    for (source, destination, port), count in syn_by_target.items():
        if port in watched_ports and count >= 8:
            alerts.append({
                "id": f"repeated-syn-{source}-{destination}-{port}", "category": "Repeated connection attempts", "severity": "info", "title": "Repeated connection attempts",
                "source_ip": source, "destination_ip": destination, "destination_port": port,
                "observed_count": count, "threshold": 8, "window_seconds": window_seconds,
                "message": f"Observed {count} TCP SYN packets from {source} to {destination}:{port} ({watched_ports[port]}) in {window_seconds} seconds. Verify whether this is expected.",
                "timestamp": now,
            })

    # Keep one alert per rule/source/target and return the most relevant first.
    unique = {}
    for alert in alerts:
        unique[alert["id"]] = alert
    severity_rank = {"warning": 0, "info": 1}
    return sorted(unique.values(), key=lambda a: (severity_rank.get(a["severity"], 2), a["title"]))[:10]


@app.route("/")
def dashboard():
    offline_error = None
    try:
        packets, protocol_counts, summary = analyze_offline_pcap()
    except Exception as error:
        app.logger.exception("Offline PCAP analysis failed")
        packets, protocol_counts = [], Counter()
        summary = {"total_packets": 0, "total_bytes": 0, "protocol_count": 0,
                   "average_packet_size": 0, "largest_packet": 0}
        offline_error = str(error)
    return render_template(
        "index.html", packets=packets, protocol_counts=dict(protocol_counts),
        summary=summary, interfaces=get_network_interfaces(),
        default_interface=WIFI_INTERFACE, offline_error=offline_error,
    )


@app.route("/api/interfaces")
def interfaces_api():
    return jsonify({"interfaces": get_network_interfaces(), "default_interface": WIFI_INTERFACE})


@app.route("/api/packets")
def offline_packets_api():
    try:
        packets, protocol_counts, summary = analyze_offline_pcap()
        return jsonify({"packets": packets, "protocol_counts": dict(protocol_counts), "summary": summary})
    except Exception as error:
        return jsonify({"error": str(error)}), 500



@app.route("/api/packet-details")
def packet_details_api():
    """Fetch full details for a selected offline or live packet."""
    source = (request.args.get("source") or "offline").lower()
    try:
        number = int(request.args.get("number", "0"))
    except (TypeError, ValueError):
        return jsonify({"error": "Packet number must be an integer."}), 400
    if number < 1:
        return jsonify({"error": "Packet number must be greater than zero."}), 400

    if source == "offline":
        try:
            if not CAPTURE_FILE.exists():
                return jsonify({"error": f"PCAP file not found: {CAPTURE_FILE}"}), 404
            packets = rdpcap(str(CAPTURE_FILE))
            if number > len(packets):
                return jsonify({"error": "Offline packet number is out of range."}), 404
            return jsonify(inspect_packet(packets[number - 1], number, "offline"))
        except Exception as error:
            app.logger.exception("Could not inspect offline packet")
            return jsonify({"error": str(error)}), 500

    if source == "live":
        with capture_lock:
            packet = live_packet_objects.get(number)
        if packet is None:
            return jsonify({"error": "This live packet is no longer available. It may have aged out of the recent packet buffer; select a newer packet."}), 404
        try:
            return jsonify(inspect_packet(packet, number, "live"))
        except Exception as error:
            app.logger.exception("Could not inspect live packet")
            return jsonify({"error": str(error)}), 500

    return jsonify({"error": "Unknown packet source. Use 'offline' or 'live'."}), 400


@app.route("/export/csv")
def export_csv():
    try:
        packets, _, _ = analyze_offline_pcap()
    except Exception as error:
        return jsonify({"error": str(error)}), 500
    output = io.StringIO()
    fields = ["number", "timestamp", "source_ip", "destination_ip", "protocol",
              "source_port", "destination_port", "size_bytes"]
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    # Export only the documented CSV columns; packet_details also contains
    # internal security-analysis fields that should not break the download.
    writer.writerows({field: packet.get(field, "") for field in fields} for packet in packets)
    return Response(output.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=packet_analysis.csv"})


@app.route("/export/report.html")
def export_analysis_report():
    """Download a self-contained HTML report for offline and current live analysis."""
    generated_at = __import__("datetime").datetime.now().astimezone().isoformat(timespec="seconds")
    offline_error = None
    try:
        offline_packets, offline_protocol_counts, offline_summary = analyze_offline_pcap()
    except Exception as error:
        offline_packets, offline_protocol_counts = [], Counter()
        offline_summary = {"total_packets": 0, "total_bytes": 0, "protocol_count": 0,
                           "average_packet_size": 0, "largest_packet": 0}
        offline_error = str(error)

    with capture_lock:
        live_history = list(live_packets)
        live_protocols = dict(live_protocol_counts)
        live_count = live_total_seen
        live_bytes = live_total_bytes
        selected_interface = live_interface or "No active capture interface"
        sniffer = live_sniffer
        try:
            capture_running = bool(sniffer is not None and sniffer.running)
        except Exception:
            capture_running = False

    analytics = build_live_analytics(live_history)
    alerts = analyze_security_alerts(live_history)

    def esc(value):
        return html_lib.escape(str(value if value is not None else ""))

    def table(headers, rows):
        head = "".join(f"<th>{esc(header)}</th>" for header in headers)
        body = "".join("<tr>" + "".join(f"<td>{esc(value)}</td>" for value in row) + "</tr>" for row in rows)
        if not body:
            body = f'<tr><td colspan="{len(headers)}">No data available.</td></tr>'
        return f"<div class='table-wrap'><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>"

    offline_protocol_rows = sorted(offline_protocol_counts.items(), key=lambda item: (-item[1], item[0]))
    live_protocol_rows = sorted(live_protocols.items(), key=lambda item: (-item[1], item[0]))
    source_rows = [(item["label"], item["count"]) for item in analytics["top_source_ips"]]
    destination_rows = [(item["label"], item["count"]) for item in analytics["top_destination_ips"]]
    port_rows = [(item["label"], item["count"]) for item in analytics["top_ports"]]
    alert_rows = [(alert["severity"].upper(), alert["title"], alert["message"],
                   __import__("datetime").datetime.fromtimestamp(float(alert["timestamp"])).astimezone().isoformat(timespec="seconds"))
                  for alert in alerts]
    latest_rows = [(p["number"], p["source_ip"], p["destination_ip"], p["protocol"],
                    p["source_port"], p["destination_port"], p["size_bytes"])
                   for p in list(reversed(live_history[-20:]))]

    error_note = f"<p class='warning'>Offline PCAP analysis issue: {esc(offline_error)}</p>" if offline_error else ""
    report = f'''<!DOCTYPE html>
<html lang="en"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Network Packet Analyzer Report</title><style>
body{{font-family:Arial,sans-serif;max-width:1100px;margin:32px auto;padding:0 18px;color:#1f2937;line-height:1.5}}
h1,h2{{color:#111827}} .muted{{color:#64748b}} .cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:12px;margin:18px 0}}
.card{{border:1px solid #dbe2ea;border-radius:10px;padding:14px;background:#f8fafc}} .card strong{{display:block;color:#64748b;font-size:12px;text-transform:uppercase}}
.card span{{font-size:24px;font-weight:700}} table{{width:100%;border-collapse:collapse;font-size:13px}} th,td{{text-align:left;padding:8px;border-bottom:1px solid #e2e8f0;overflow-wrap:anywhere}} th{{background:#f1f5f9}}
.table-wrap{{overflow-x:auto;margin:10px 0 24px}} section{{margin-top:28px}} .warning{{background:#fff7ed;padding:12px;border-left:4px solid #ea580c}} .alert{{border-left:4px solid #d97706;padding:8px 12px;background:#fffbeb;margin:8px 0}}
@media print{{body{{margin:0;max-width:none}} .no-print{{display:none}}}}
</style></head><body>
<h1>Network Packet Analyzer Report</h1><p class="muted">Generated: {esc(generated_at)} · Live capture: {"Running" if capture_running else "Stopped"}</p>
<p class="no-print">Use your browser's Print command to save this report as PDF.</p>{error_note}
<section><h2>Offline PCAP Summary</h2><div class="cards">
<div class="card"><strong>Total packets</strong><span>{offline_summary['total_packets']}</span></div>
<div class="card"><strong>Total bytes</strong><span>{offline_summary['total_bytes']:,}</span></div>
<div class="card"><strong>Protocols</strong><span>{offline_summary['protocol_count']}</span></div>
<div class="card"><strong>Average packet</strong><span>{offline_summary['average_packet_size']} B</span></div>
<div class="card"><strong>Largest packet</strong><span>{offline_summary['largest_packet']} B</span></div></div>
<h3>Offline protocol distribution</h3>{table(["Protocol", "Packets"], offline_protocol_rows)}</section>
<section><h2>Live Capture Summary</h2><div class="cards">
<div class="card"><strong>Total captured packets</strong><span>{live_count:,}</span></div>
<div class="card"><strong>Total captured bytes</strong><span>{live_bytes:,}</span></div>
<div class="card"><strong>Protocols observed</strong><span>{len(live_protocols)}</span></div>
<div class="card"><strong>Retained packet history</strong><span>{len(live_history):,}</span></div></div>
<p><strong>Interface:</strong> {esc(selected_interface)} · <strong>Capture state:</strong> {"Running" if capture_running else "Stopped"}</p>
<h3>Live protocol distribution</h3>{table(["Protocol", "Packets"], live_protocol_rows)}
<h3>Top source IP addresses (retained history)</h3>{table(["Source IP", "Packets"], source_rows)}
<h3>Top destination IP addresses (retained history)</h3>{table(["Destination IP", "Packets"], destination_rows)}
<h3>Most active ports (retained history)</h3>{table(["Port / Protocol", "Observations"], port_rows)}
<h3>Latest retained packets</h3>{table(["No.", "Source IP", "Destination IP", "Protocol", "Source Port", "Destination Port", "Bytes"], latest_rows)}</section>
<section><h2>Security Monitoring</h2><p class="muted">Heuristic indicators based on the recent retained live packet history. Alerts are not proof of malicious activity, and no alerts do not guarantee safety.</p>
{''.join(f'<div class="alert"><strong>{esc(alert["severity"].upper())} — {esc(alert["title"])}</strong><p>{esc(alert["message"])}</p></div>' for alert in alerts) if alerts else '<p>No heuristic alerts are present in the current evaluation window.</p>'}</section>
<footer class="muted"><p>Report generated locally by Network Packet Analyzer.</p></footer></body></html>'''
    return Response(report, mimetype="text/html; charset=utf-8",
                    headers={"Content-Disposition": "attachment; filename=network_analysis_report.html"})


@app.route("/export/security.csv")
def export_security_csv():
    """Download a detailed, spreadsheet-friendly CSV of current heuristic alerts."""
    with capture_lock:
        history = list(live_packets)
        selected_interface = live_interface or "No active capture interface"
    alerts = analyze_security_alerts(history)
    generated_at = __import__("datetime").datetime.now().astimezone().isoformat(timespec="seconds")
    output = io.StringIO(newline="")
    fields = [
        "generated_at_local", "alert_time_local", "severity", "category", "alert_id",
        "title", "source_ip", "destination_ip", "destination_port", "observed_count",
        "threshold", "window_seconds", "capture_interface", "message", "interpretation_note",
    ]
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for alert in alerts:
        writer.writerow({
            "generated_at_local": generated_at,
            "alert_time_local": __import__("datetime").datetime.fromtimestamp(float(alert.get("timestamp", time.time()))).astimezone().isoformat(timespec="seconds"),
            "severity": str(alert.get("severity", "info")).upper(),
            "category": alert.get("category", alert.get("title", "General")),
            "alert_id": alert.get("id", ""),
            "title": alert.get("title", ""),
            "source_ip": alert.get("source_ip", ""),
            "destination_ip": alert.get("destination_ip", ""),
            "destination_port": alert.get("destination_port", ""),
            "observed_count": alert.get("observed_count", ""),
            "threshold": alert.get("threshold", ""),
            "window_seconds": alert.get("window_seconds", ""),
            "capture_interface": selected_interface,
            "message": alert.get("message", ""),
            "interpretation_note": "Heuristic indicator only; not proof of malicious activity.",
        })
    # Excel recognizes UTF-8 BOM and correctly opens non-ASCII text in CSV.
    return Response("\ufeff" + output.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": "attachment; filename=security_alerts_detailed.csv"})


@app.route("/export/security.xlsx")
def export_security_xlsx():
    """Download a formatted Excel security workbook with summary charts and detailed alerts."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.chart import PieChart, BarChart, Reference
        from openpyxl.chart.label import DataLabelList
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({
            "error": "Excel export requires openpyxl. Install it in your project environment with: python -m pip install openpyxl"
        }), 500

    with capture_lock:
        history = list(live_packets)
        selected_interface = live_interface or "No active capture interface"
    alerts = analyze_security_alerts(history)
    generated_at = __import__("datetime").datetime.now().astimezone().isoformat(timespec="seconds")

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Security Dashboard"
    detail = workbook.create_sheet("Alert Details")

    navy = "17365D"
    blue = "D9EAF7"
    white = "FFFFFF"
    red_fill = PatternFill("solid", fgColor="FCE4D6")
    yellow_fill = PatternFill("solid", fgColor="FFF2CC")
    green_fill = PatternFill("solid", fgColor="E2F0D9")
    thin_gray = Side(style="thin", color="D9E1F2")

    summary.merge_cells("A1:H1")
    summary["A1"] = "NETWORK PACKET ANALYZER — SECURITY ALERT REPORT"
    summary["A1"].font = Font(bold=True, size=16, color=white)
    summary["A1"].fill = PatternFill("solid", fgColor=navy)
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 30
    summary["A3"] = "Generated (local time)"
    summary["B3"] = generated_at
    summary["A4"] = "Capture interface"
    summary["B4"] = selected_interface
    summary["A5"] = "Retained packets analyzed"
    summary["B5"] = len(history)
    summary["A6"] = "Alerts in current evaluation"
    summary["B6"] = len(alerts)
    summary["A7"] = "Important note"
    summary["B7"] = "Heuristic indicators only; not proof of malicious activity. No alerts does not guarantee safety."
    summary.merge_cells("B7:H7")
    summary["B7"].alignment = Alignment(wrap_text=True, vertical="top")
    summary.row_dimensions[7].height = 36
    for cell in (summary["A3"], summary["A4"], summary["A5"], summary["A6"], summary["A7"]):
        cell.font = Font(bold=True, color=navy)

    severity_counts = Counter(str(alert.get("severity", "info")).upper() for alert in alerts)
    category_counts = Counter(str(alert.get("category", alert.get("title", "Other"))) for alert in alerts)
    summary["A10"] = "Severity"
    summary["B10"] = "Alert count"
    summary["D10"] = "Alert category"
    summary["E10"] = "Alert count"
    for cell in (summary["A10"], summary["B10"], summary["D10"], summary["E10"]):
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)

    severity_rows = [("WARNING", severity_counts.get("WARNING", 0)), ("INFO", severity_counts.get("INFO", 0))]
    for row_num, (severity, count) in enumerate(severity_rows, start=11):
        summary.cell(row_num, 1, severity)
        summary.cell(row_num, 2, count)
    categories = sorted(category_counts.items(), key=lambda item: (-item[1], item[0])) or [("No alerts", 0)]
    for row_num, (category, count) in enumerate(categories, start=11):
        summary.cell(row_num, 4, category)
        summary.cell(row_num, 5, count)

    pie = PieChart()
    pie.title = "Alerts by Severity"
    pie.height = 7
    pie.width = 10
    pie.add_data(Reference(summary, min_col=2, min_row=10, max_row=12), titles_from_data=True)
    pie.set_categories(Reference(summary, min_col=1, min_row=11, max_row=12))
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showLeaderLines = True
    summary.add_chart(pie, "A15")

    bar = BarChart()
    bar.type = "bar"
    bar.style = 10
    bar.title = "Alerts by Category"
    bar.y_axis.title = "Category"
    bar.x_axis.title = "Alert count"
    bar.height = 8
    bar.width = 13
    bar.add_data(Reference(summary, min_col=5, min_row=10, max_row=10 + len(categories)), titles_from_data=True)
    bar.set_categories(Reference(summary, min_col=4, min_row=11, max_row=10 + len(categories)))
    bar.legend = None
    summary.add_chart(bar, "D15")

    headers = [
        "Generated At (Local)", "Alert Time (Local)", "Severity", "Category", "Alert ID",
        "Title", "Source IP", "Destination IP", "Destination Port", "Observed Count",
        "Threshold", "Window (seconds)", "Capture Interface", "Message", "Interpretation Note",
    ]
    detail.append(headers)
    for alert in alerts:
        local_time = __import__("datetime").datetime.fromtimestamp(float(alert.get("timestamp", time.time()))).astimezone().isoformat(timespec="seconds")
        detail.append([
            generated_at, local_time, str(alert.get("severity", "info")).upper(),
            alert.get("category", alert.get("title", "General")), alert.get("id", ""),
            alert.get("title", ""), alert.get("source_ip", ""), alert.get("destination_ip", ""),
            alert.get("destination_port", ""), alert.get("observed_count", ""),
            alert.get("threshold", ""), alert.get("window_seconds", ""), selected_interface,
            alert.get("message", ""), "Heuristic indicator only; not proof of malicious activity.",
        ])
    if not alerts:
        detail.append([generated_at, "", "NONE", "No current alerts", "", "No alerts detected in the current evaluation window", "", "", "", 0, "", 30, selected_interface, "No heuristic alerts were generated from the retained recent packet history.", "No alerts does not guarantee that the network is safe."])

    for cell in detail[1]:
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    detail.row_dimensions[1].height = 32
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    widths = [24, 24, 13, 28, 36, 30, 20, 30, 18, 17, 12, 17, 42, 85, 62]
    for col_idx, width in enumerate(widths, start=1):
        detail.column_dimensions[get_column_letter(col_idx)].width = width
    for row in detail.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        severity = str(row[2].value or "").upper()
        fill = red_fill if severity == "WARNING" else yellow_fill if severity == "INFO" else green_fill
        row[2].fill = fill
        row[2].font = Font(bold=True)
    if alerts:
        table_ref = f"A1:O{detail.max_row}"
        excel_table = Table(displayName="SecurityAlerts", ref=table_ref)
        excel_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        detail.add_table(excel_table)

    for col, width in {"A": 26, "B": 22, "D": 32, "E": 18, "H": 20}.items():
        summary.column_dimensions[col].width = width
    summary.column_dimensions["C"].width = 3
    summary.column_dimensions["F"].width = 16
    summary.column_dimensions["G"].width = 16
    summary.column_dimensions["H"].width = 16
    summary.freeze_panes = "A10"

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return Response(output.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=security_alerts_report.xlsx"})


@app.route("/api/live/start", methods=["POST"])
def start_live_capture():
    global live_sniffer, live_interface, live_error, live_total_bytes, live_total_seen
    data = request.get_json(silent=True) or {}
    interface = data.get("interface") or WIFI_INTERFACE
    valid_interfaces = {item["name"] for item in get_network_interfaces()}
    if interface not in valid_interfaces:
        return jsonify({"error": "Invalid interface. Refresh the page and select a listed adapter."}), 400

    with capture_lock:
        if live_sniffer is not None:
            try:
                if live_sniffer.running:
                    return jsonify({"error": "Capture is already running.", "interface": live_interface}), 409
            except Exception:
                pass
            live_sniffer = None
            live_interface = None
        live_packets.clear()
        live_packet_objects.clear()
        live_protocol_counts.clear()
        live_total_bytes = 0
        live_total_seen = 0
        live_error = None
        sniffer = AsyncSniffer(iface=interface, prn=live_packet_callback, store=False)
        live_sniffer = sniffer
        live_interface = interface

    try:
        sniffer.start()
        return jsonify({"status": "started", "running": True, "interface": interface,
                        "message": "Capture started. Generate network traffic to see packets."})
    except Exception as error:
        with capture_lock:
            if live_sniffer is sniffer:
                live_sniffer = None
                live_interface = None
            live_error = str(error)
        app.logger.exception("Live capture failed to start")
        return jsonify({"error": str(error), "hint": "Try the confirmed Wi-Fi interface and run Flask as Administrator if Npcap permissions are required."}), 500


@app.route("/api/live/stop", methods=["POST"])
def stop_live_capture():
    global live_sniffer, live_interface, live_error
    with capture_lock:
        sniffer = live_sniffer
    if sniffer is None:
        return jsonify({"status": "stopped", "message": "No capture session is running."})
    try:
        if sniffer.running:
            sniffer.stop(join=True)
        with capture_lock:
            if live_sniffer is sniffer:
                live_sniffer = None
                live_interface = None
        return jsonify({"status": "stopped", "message": "Live capture stopped. Captured statistics are retained."})
    except Exception as error:
        with capture_lock:
            live_error = str(error)
        app.logger.exception("Could not stop live capture")
        return jsonify({"error": str(error)}), 500


@app.route("/api/live/status")
def live_status():
    with capture_lock:
        sniffer = live_sniffer
        try:
            running = bool(sniffer is not None and sniffer.running)
        except Exception:
            running = False
        recent_history = list(live_packets)
        analytics = build_live_analytics(recent_history)
        return jsonify({
            "running": running, "interface": live_interface,
            "total_packets": live_total_seen, "total_bytes": live_total_bytes,
            "protocol_counts": dict(live_protocol_counts),
            "recent_packets": recent_history, "analytics": analytics,
            "security_alerts": analyze_security_alerts(recent_history),
            "error": live_error,
        })


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False, threaded=True)
